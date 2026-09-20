# -*- coding: utf-8 -*-
"""
单元测试
--------
覆盖核心模块：打分引擎、条件筛选、存储层、Excel 导出。
运行：python -m pytest tests/ -v
"""
import sys
from pathlib import Path

# 把项目根目录加入 sys.path，便于直接运行
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest

from config import OUTPUT_COLUMNS
from core.models import Job
from core.matcher import JobMatcher
from core.filters import JobFilter, parse_salary_to_monthly
from core.storage import JobStorage
from core.exporter import ExcelExporter, FIELD_TO_HEADER


def make_row(**kwargs) -> dict:
    """
    构造一条模拟 storage.query_jobs() 返回的记录（英文 key）。

    默认值模拟一条「厦门、本科可投、匹配分较高」的普通岗位，
    测试里只想改某一两个字段时，直接覆盖即可。
    """
    row = {
        "id": 1,
        "company": "厦门某科技公司",
        "title": "算法工程师",
        "city": "福建省厦门市思明区",
        "salary": "8000-12000",
        "education": "本科",
        "major_requirement": "数学类,信息与计算科学",
        "apply_method": "https://example.com",
        "deadline": "2026-10-01",
        "source": "厦门大学就业信息网",
        "url": "https://example.com",
        "match_score": 25,
        "match_level": "S",
        "match_label": "极高匹配",
        "hit_keywords": "信息与计算科学,数学类",
        "crawl_time": "2026-09-18 12:00:00",
    }
    row.update(kwargs)
    return row


# ===================================================================
# 数据模型测试
# ===================================================================
class TestJobModel:

    def test_is_valid(self):
        """公司或岗位名为空时视为无效"""
        assert Job(company="A", title="B").is_valid() is True
        assert Job(company="", title="B").is_valid() is False
        assert Job(company="A", title="").is_valid() is False

    def test_to_dict_converts_keywords(self):
        """字典转换时应把关键词列表转为逗号串"""
        job = Job(company="A", title="B", hit_keywords=["数学类", "统计"])
        assert job.to_dict()["hit_keywords"] == "数学类,统计"


# ===================================================================
# 打分引擎测试
# ===================================================================
class TestJobMatcher:

    @pytest.fixture
    def matcher(self):
        return JobMatcher()

    def test_exact_major_match_scores_highest(self, matcher):
        """直接点名专业应得最高分"""
        job = Job(
            company="网宿科技", title="算法工程师",
            major_requirement="数学类（数学与应用数学、信息与计算科学）",
        )
        score, hits = matcher.score(job)
        assert score >= 25, f"直接点名专业的岗位应达 S 级，实际 {score} 分"
        assert "信息与计算科学" in hits
        assert matcher.level_of(score)[0] == "S"

    def test_statistics_major_match(self, matcher):
        """数理统计类要求应获高分"""
        job = Job(
            company="建设银行", title="科技类专项人才",
            major_requirement="重点招收计算机、人工智能、数理统计及其他理工类专业",
        )
        score, hits = matcher.score(job)
        assert score >= 16, f"数理统计类应达 A 级，实际 {score} 分"
        assert "数理统计" in hits

    def test_negative_keyword_deducts(self, matcher):
        """销售类岗位应被扣分"""
        job = Job(
            company="某公司", title="销售代表",
            major_requirement="专业不限",
        )
        score, hits = matcher.score(job)
        assert score == 0, f"销售岗应被扣至 0 分，实际 {score}"
        assert any("扣分" in h for h in hits)

    def test_empty_job_scores_zero(self, matcher):
        """空岗位得 0 分"""
        score, hits = matcher.score(Job())
        assert score == 0
        assert hits == []

    def test_score_jobs_sorts_descending(self, matcher):
        """批量打分后应按分数降序"""
        jobs = [
            Job(company="C", title="岗位C", major_requirement="专业不限"),
            Job(company="A", title="岗位A",
                major_requirement="信息与计算科学"),
            Job(company="B", title="岗位B", major_requirement="数学类"),
        ]
        result = matcher.score_jobs(jobs)
        scores = [j.match_score for j in result]
        assert scores == sorted(scores, reverse=True), "应按分数降序排列"
        assert result[0].company == "A"

    def test_level_mapping(self, matcher):
        """等级映射边界检查"""
        assert matcher.level_of(30)[0] == "S"
        assert matcher.level_of(25)[0] == "S"
        assert matcher.level_of(16)[0] == "A"
        assert matcher.level_of(8)[0] == "B"
        assert matcher.level_of(0)[0] == "C"


# ===================================================================
# 条件筛选测试
# ===================================================================
class TestSalaryParsing:
    """薪资文本解析"""

    @pytest.mark.parametrize("text,expected", [
        ("7000-10000", 7000),
        ("8000", 8000),
        ("8000元/月", 8000),
        ("10-15万/年", 8333),
        ("年薪20万", 16666),
        ("面议", 0),
        ("", 0),
        (None, 0),
        ("待遇丰厚", 0),
    ])
    def test_parse(self, text, expected):
        assert parse_salary_to_monthly(text) == expected


class TestJobFilter:
    """
    用户条件筛选。

    这一组测试锁定的是一整套「条件 -> 结果」的契约。之所以单独成一个类，
    是因为筛选和打分是两件事：打分管「对口不对口」，筛分管「我要不要看」。
    """

    def test_empty_filter_returns_all(self):
        """不设任何条件时不应丢数据"""
        rows = [make_row(id=1), make_row(id=2, company="B"), make_row(id=3, company="C")]
        # bachelor_only 默认开启，这里构造的三条都是本科可投
        result = JobFilter().apply(rows)
        assert len(result) == 3

    # ---------------- 城市 ----------------
    def test_filter_by_city(self):
        """城市条件应只保留命中的城市"""
        rows = [
            make_row(id=1, city="福建省厦门市"),
            make_row(id=2, city="福建省福州市"),
            make_row(id=3, city="湖南省长沙市"),
        ]
        result = JobFilter(city=["厦门"]).apply(rows)
        assert len(result) == 1
        assert result[0]["city"] == "福建省厦门市"

    def test_filter_by_multiple_cities(self):
        """多城市应取并集"""
        rows = [
            make_row(id=1, city="福建省厦门市"),
            make_row(id=2, city="福建省福州市"),
            make_row(id=3, city="湖南省长沙市"),
        ]
        result = JobFilter(city=["厦门", "福州"]).apply(rows)
        assert len(result) == 2

    def test_city_priority(self):
        """城市优先级：目标城市 < 同省 < 其他/未知"""
        assert JobFilter.city_priority("福建省厦门市", "厦门") == 0
        assert JobFilter.city_priority("福建省福州市", "厦门") == 1
        assert JobFilter.city_priority("湖南省长沙市", "厦门") == 2
        assert JobFilter.city_priority("", "厦门") == 2

    def test_sort_puts_target_city_first(self):
        """
        ★ 回归测试：目标城市优先必须真正生效。

        背景（实际调试中发现的失效 Bug）：
            原先在流水线里先按城市排了一次序，紧接着 score_jobs() 又按分数
            重排了一次，城市排序结果被直接覆盖 —— 等于这段代码从来没生效过。
            现在排序和筛选统一在 JobFilter.apply() 里做，才真正起作用。
        """
        rows = [
            make_row(id=1, city="湖南省长沙市", match_score=90),
            make_row(id=2, city="福建省厦门市", match_score=10),
        ]
        result = JobFilter().apply(rows)
        assert result[0]["city"] == "福建省厦门市", "厦门岗位必须排在最前（哪怕分数更低）"
        assert result[1]["city"] == "湖南省长沙市"

    def test_sort_by_score_within_same_city(self):
        """同一城市优先级内，仍按匹配分降序"""
        rows = [
            make_row(id=1, city="福建省厦门市", match_score=10),
            make_row(id=2, city="福建省厦门市", match_score=30),
        ]
        result = JobFilter().apply(rows)
        assert [r["match_score"] for r in result] == [30, 10]

    # ---------------- 分数 / 等级 ----------------
    def test_min_score(self):
        rows = [
            make_row(id=1, match_score=30),
            make_row(id=2, match_score=10),
            make_row(id=3, match_score=20),
        ]
        result = JobFilter(min_score=20).apply(rows)
        assert len(result) == 2
        assert all(r["match_score"] >= 20 for r in result)

    def test_levels_accept_multiple(self):
        """等级支持多选（S,A）"""
        rows = [
            make_row(id=1, match_level="S", match_score=30),
            make_row(id=2, match_level="A", match_score=20),
            make_row(id=3, match_level="C", match_score=2),
        ]
        result = JobFilter(levels=["S", "A"]).apply(rows)
        assert len(result) == 2
        assert {r["match_level"] for r in result} == {"S", "A"}

    def test_levels_case_insensitive(self):
        """等级输入小写也应能用"""
        rows = [make_row(id=1, match_level="S", match_score=30)]
        assert len(JobFilter(levels=["s"]).apply(rows)) == 1

    # ---------------- 关键词 ----------------
    def test_keyword_matches_title(self):
        rows = [make_row(id=1, title="算法工程师"),
                make_row(id=2, title="市场营销专员", match_score=1)]
        result = JobFilter(keyword="算法").apply(rows)
        assert len(result) == 1

    def test_keyword_matches_company(self):
        rows = [make_row(id=1, company="厦门点触科技股份有限公司"),
                make_row(id=2, company="其他公司")]
        result = JobFilter(keyword="点触").apply(rows)
        assert len(result) == 1

    def test_major_contains(self):
        rows = [
            make_row(id=1, major_requirement="数学类,信息与计算科学"),
            make_row(id=2, major_requirement="专业不限"),
        ]
        result = JobFilter(major="数学").apply(rows)
        assert len(result) == 1

    def test_exclude_keywords(self):
        rows = [
            make_row(id=1, title="销售代表", match_score=5),
            make_row(id=2, title="算法工程师", match_score=5),
        ]
        result = JobFilter(exclude=["销售", "客服"]).apply(rows)
        assert len(result) == 1
        assert result[0]["title"] == "算法工程师"

    # ---------------- 学历 ----------------
    def test_bachelor_only_removes_master_only(self):
        """默认应剔除仅招硕博的岗位"""
        rows = [
            make_row(id=1, education="本科及以上"),
            make_row(id=2, education="全日制硕士/博士"),
            make_row(id=3, education="硕士,本科"),
        ]
        result = JobFilter().apply(rows)
        assert len(result) == 2
        assert all(r["id"] != 2 for r in result)

    def test_all_master_keeps_them(self):
        """显式关闭本科过滤后，硕博岗位应保留"""
        rows = [make_row(id=1, education="全日制硕士/博士")]
        result = JobFilter(bachelor_only=False).apply(rows)
        assert len(result) == 1

    def test_education_unknown_is_kept(self):
        """学历为空的岗位视为未知，不应被 --education 误杀"""
        rows = [make_row(id=1, education="")]
        assert len(JobFilter(education="本科").apply(rows)) == 1

    # ---------------- 薪资 ----------------
    def test_salary_min_filters_low(self):
        rows = [
            make_row(id=1, salary="6000-8000"),
            make_row(id=2, salary="12000-15000"),
        ]
        result = JobFilter(salary_min=8000).apply(rows)
        assert len(result) == 1
        assert result[0]["salary"] == "12000-15000"

    def test_salary_unknown_is_kept(self):
        """薪资未标注（面议）的岗位不应因为信息缺失被剔除"""
        rows = [make_row(id=1, salary="面议"), make_row(id=2, salary="")]
        assert len(JobFilter(salary_min=20000).apply(rows)) == 2

    # ---------------- 条数 ----------------
    def test_top_limit(self):
        rows = [make_row(id=i, company=f"C{i}", match_score=i) for i in range(1, 6)]
        result = JobFilter(top=2).apply(rows)
        assert len(result) == 2

    def test_top_applied_after_sort(self):
        """top 必须作用在排序之后，否则会截掉分数最高的岗位"""
        rows = [
            make_row(id=1, company="低分", match_score=1),
            make_row(id=2, company="高分", match_score=50),
        ]
        result = JobFilter(top=1).apply(rows)
        assert result[0]["company"] == "高分"

    # ---------------- 组合 ----------------
    def test_conditions_are_combined_with_and(self):
        """多个条件之间是 AND 关系"""
        rows = [
            make_row(id=1, city="福建省厦门市", match_score=30, match_level="S"),
            make_row(id=2, city="福建省厦门市", match_score=5, match_level="C"),
            make_row(id=3, city="湖南省长沙市", match_score=30, match_level="S"),
        ]
        result = JobFilter(city=["厦门"], min_score=20, levels=["S"]).apply(rows)
        assert len(result) == 1
        assert result[0]["id"] == 1

    def test_does_not_mutate_input(self):
        """apply() 不应修改传入的列表"""
        rows = [make_row(id=1, city="湖南省长沙市", match_score=1),
                make_row(id=2, city="福建省厦门市", match_score=2)]
        snapshot = [dict(r) for r in rows]
        JobFilter().apply(rows)
        assert rows == snapshot, "apply() 必须是无副作用的"

    # ---------------- 方案保存 / 载入 ----------------
    def test_save_and_load_roundtrip(self, tmp_path):
        """一套条件存成 JSON 后应能原样载入"""
        original = JobFilter(
            city=["厦门", "福州"], min_score=16, levels=["S", "A"],
            keyword="算法", major="数学", education="本科",
            salary_min=8000, exclude=["销售"], bachelor_only=False, top=20,
        )
        path = original.save(tmp_path / "plan.json")
        loaded = JobFilter.load(path)
        assert loaded.to_dict() == original.to_dict()

    def test_load_missing_file_returns_empty_filter(self, tmp_path):
        """方案文件不存在时应返回空条件，而不是抛异常"""
        loaded = JobFilter.load(tmp_path / "not_exist.json")
        # 「空条件」指的是没有城市/分数/等级等限制，
        # 但 bachelor_only 属于默认开启的行为开关，不受方案文件影响
        assert loaded.city == []
        assert loaded.min_score == 0
        assert loaded.levels == []
        assert loaded.bachelor_only is True
        assert loaded.describe() == ["仅本科可投"]

    def test_from_dict_ignores_unknown_fields(self):
        """方案文件里多写了字段也不应导致崩溃"""
        loaded = JobFilter.from_dict({"city": ["厦门"], "不存在的字段": 1})
        assert loaded.city == ["厦门"]

    # ---------------- 描述 ----------------
    def test_describe_lists_active_conditions(self):
        desc = " | ".join(JobFilter(city=["厦门"], min_score=16).describe())
        assert "厦门" in desc
        assert "16" in desc


# ===================================================================
# 存储层测试
# ===================================================================
class TestJobStorage:

    @pytest.fixture
    def storage(self, tmp_path):
        """使用临时数据库，避免污染真实数据"""
        return JobStorage(db_path=tmp_path / "test.db")

    def test_save_and_count(self, storage):
        """写入后应能统计到数量"""
        jobs = [
            Job(company="A", title="岗位A"),
            Job(company="B", title="岗位B"),
        ]
        inserted = storage.save_jobs(jobs)
        assert inserted == 2
        assert storage.count() == 2

    def test_duplicate_is_ignored(self, storage):
        """重复岗位应被忽略"""
        job = Job(company="A", title="岗位A")
        storage.save_jobs([job])
        storage.save_jobs([job])
        assert storage.count() == 1, "相同公司+岗位应只保留一条"

    def test_query_all_sorted_by_score(self, storage):
        """不传条件时应返回全部，且按分数降序"""
        storage.save_jobs([
            Job(company="A", title="a", match_score=10),
            Job(company="B", title="b", match_score=30),
            Job(company="C", title="c", match_score=20),
        ])
        result = storage.query_jobs()
        assert [r["match_score"] for r in result] == [30, 20, 10]

    def test_query_applies_filter(self, storage):
        """
        ★ 回归测试：query_jobs 必须真的调用传入的 JobFilter。

        背景：JobFilter 的前身是 query_jobs 上的 city/min_score/level 三个
        关键字参数，功能写好了但全项目没有任何调用方——能力造出来没接线。
        这个测试保证「筛选条件」和「取数」之间的连接不会再次断掉。
        """
        storage.save_jobs([
            Job(company="A", title="a", city="福建省厦门市", match_score=30),
            Job(company="B", title="b", city="福建省厦门市", match_score=5),
            Job(company="C", title="c", city="湖南省长沙市", match_score=30),
        ])
        result = storage.query_jobs(job_filter=JobFilter(city=["厦门"], min_score=20))
        assert len(result) == 1
        assert result[0]["company"] == "A"

    def test_query_limit_applied_after_filter(self, storage):
        """limit 必须在筛选之后生效"""
        storage.save_jobs([
            Job(company="A", title="a", match_score=30),
            Job(company="B", title="b", match_score=20),
            Job(company="C", title="c", match_score=10),
        ])
        result = storage.query_jobs(job_filter=JobFilter(), limit=2)
        assert [r["company"] for r in result] == ["A", "B"]

    def test_count_by_level(self, storage):
        """按等级统计"""
        jobs = [
            Job(company="A", title="a", match_level="S"),
            Job(company="B", title="b", match_level="S"),
            Job(company="C", title="c", match_level="A"),
        ]
        storage.save_jobs(jobs)
        dist = storage.count_by_level()
        assert dist.get("S") == 2
        assert dist.get("A") == 1

    def test_count_by_city(self, storage):
        """按城市统计"""
        storage.save_jobs([
            Job(company="A", title="a", city="福建省厦门市"),
            Job(company="B", title="b", city="福建省厦门市"),
            Job(company="C", title="c", city="湖南省长沙市"),
        ])
        dist = dict(storage.count_by_city())
        assert dist["福建省厦门市"] == 2
        assert dist["湖南省长沙市"] == 1


# ===================================================================
# 导出层测试
# ===================================================================
class TestExcelExporter:
    """
    回归测试：导出的 Excel 必须真的有数据。

    背景（实际调试中发现的 Bug）：
        storage.query_jobs() 返回的字典 key 是**英文**（company/title/...），
        而 config.OUTPUT_COLUMNS 是**中文**表头（公司/岗位/...）。
        两者对不上时，pandas 会把每个中文列都补成空字符串，
        导出的 Excel 只有表头、数据行全空——而且不会报任何错。
        这个测试就是用来锁死这个静默失败的。
    """

    @pytest.fixture
    def sample_rows(self):
        """模拟 storage.query_jobs() 的真实返回结构（英文 key）"""
        return [
            {
                "id": 1,
                "company": "厦门点触科技股份有限公司",
                "title": "27届校招-游戏数值策划",
                "city": "福建省厦门市思明区",
                "salary": "7000-10000",
                "education": "本科",
                "major_requirement": "数学类,信息与计算科学",
                "apply_method": "https://jy.xmu.edu.cn/job/view/id/2397816",
                "deadline": "2026-09-03",
                "source": "厦门大学就业信息网",
                "url": "https://jy.xmu.edu.cn/job/view/id/2397816",
                "match_score": 77,
                "match_level": "S",
                "match_label": "极高匹配",
                "hit_keywords": "信息与计算科学,数学类",
                "crawl_time": "2026-09-17 19:34:09",
            },
            {
                "id": 2,
                "company": "网宿科技",
                "title": "大模型训练工程师",
                "city": "福建省厦门市",
                "salary": "面议",
                "education": "本科",
                "major_requirement": "数学类,计算机类",
                "apply_method": "https://example.com",
                "deadline": "",
                "source": "厦门大学就业信息网",
                "url": "https://example.com",
                "match_score": 36,
                "match_level": "S",
                "match_label": "极高匹配",
                "hit_keywords": "数学类,计算机类",
                "crawl_time": "2026-09-17 19:34:10",
            },
        ]

    def _read_sheet(self, path):
        import openpyxl
        return openpyxl.load_workbook(path).active

    def test_export_writes_actual_data(self, tmp_path, sample_rows):
        """★ 核心回归：数据行不能为空"""
        exporter = ExcelExporter(output_dir=tmp_path)
        path = exporter.export(sample_rows, filename="test.xlsx")

        ws = self._read_sheet(path)
        assert ws.max_row == 3, "1 行表头 + 2 行数据"

        # 表头必须是中文
        headers = [c.value for c in ws[1]]
        assert headers == OUTPUT_COLUMNS, f"表头顺序/内容不符：{headers}"

        # 数据行必须有值 —— 这是 Bug 的直接断言点
        row2 = [c.value for c in ws[2]]
        assert row2[0] == "S", "匹配等级应为 S"
        assert row2[1] == 77, "匹配分应为 77"
        assert "点触科技" in str(row2[3]), "公司名应正确写入"
        assert "游戏数值策划" in str(row2[4]), "岗位名应正确写入"
        assert "厦门" in str(row2[5]), "城市应正确写入"

        # 所有数据列都不应整列为空
        for col_idx in range(len(OUTPUT_COLUMNS)):
            values = [ws.cell(row=r, column=col_idx + 1).value
                      for r in range(2, ws.max_row + 1)]
            assert any(v not in (None, "") for v in values), \
                f"列「{OUTPUT_COLUMNS[col_idx]}」整列为空，疑似字段映射错位"

    def test_header_mapping_covers_all_fields(self, sample_rows):
        """映射表必须覆盖 OUTPUT_COLUMNS 里的每一个中文表头"""
        mapped_headers = set(FIELD_TO_HEADER.values())
        for col in OUTPUT_COLUMNS:
            assert col in mapped_headers, f"表头「{col}」没有对应的英文字段名"

    def test_export_rejects_empty_list(self, tmp_path):
        """空列表应显式报错，而不是导出空表"""
        exporter = ExcelExporter(output_dir=tmp_path)
        with pytest.raises(ValueError, match="岗位列表为空"):
            exporter.export([])

    def test_export_rejects_unknown_schema(self, tmp_path):
        """
        字段名完全无法识别时应报错。
        这是防止「静默导出空表」的第二道防线。
        """
        exporter = ExcelExporter(output_dir=tmp_path)
        bogus = [{"foo": "bar", "baz": 1}]
        with pytest.raises(ValueError, match="无法识别"):
            exporter.export(bogus, filename="bogus.xlsx")


class TestDependencyCheck:
    """
    回归测试：缺依赖时必须给出「可执行」的提示，而不是裸 traceback。

    背景（实际调试中发现的 Bug）：
        本机 python 指向系统 Python（D:\\Python313），它没装 pandas，
        而且连 pip 都没有。用户敲 python main.py 时，terminal 只回一段
        ModuleNotFoundError 堆栈 —— 既不知道正在用哪个解释器，也不知道怎么修，
        看起来像"项目坏了"。
        main.check_dependencies() 就是用来把这段堆栈换成可执行的指引。
    """

    def test_required_deps_are_well_formed(self):
        """清单必须是 (import名, pip包名, 用途) 三元组，且覆盖全部关键包"""
        import main

        assert main.REQUIRED_DEPS, "依赖清单不能为空"
        for item in main.REQUIRED_DEPS:
            assert len(item) == 3, f"清单项格式错误：{item}"
            module_name, package_name, purpose = item
            assert module_name and package_name and purpose

        modules = {m for m, _, _ in main.REQUIRED_DEPS}
        for must_have in ("pandas", "openpyxl", "requests", "bs4", "lxml"):
            assert must_have in modules, f"依赖清单漏了 {must_have}"

    def test_passes_silently_when_all_installed(self, capsys):
        """当前 venv 依赖齐全，自检应静默通过：不打印任何东西、不退出"""
        import main

        main.check_dependencies()

        assert capsys.readouterr().out == ""

    def test_exits_with_actionable_message_when_missing(self, monkeypatch, capsys):
        """缺包时必须：退出码 1 + 打印当前解释器 + 打印可直接照抄的修复命令"""
        import main

        monkeypatch.setattr(
            main,
            "REQUIRED_DEPS",
            [("definitely_not_a_real_module_xyz", "fake-pkg", "测试用")],
        )

        with pytest.raises(SystemExit) as exc_info:
            main.check_dependencies()

        assert exc_info.value.code == 1

        out = capsys.readouterr().out
        assert "缺少运行依赖" in out
        assert sys.executable in out, "必须告诉用户当前用的是哪个解释器"
        assert "fake-pkg" in out, "必须给出 pip 包名，否则用户无从搜索"
        assert "run.bat" in out, "必须给出可直接照抄的修复命令"
        assert "运行说明.md" in out, "必须指向详细文档"



if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-v", "--tb=short"]))
